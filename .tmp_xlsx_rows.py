import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')
p = r'd:\Users\HP\Downloads\xuanshiai problem\2026_10_27_用户端测试问题\用户端测试问题.xlsx'
wb = openpyxl.load_workbook(p, data_only=True)
ws = wb.active
module = ''
for i, row in enumerate(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True), start=1):
    vals = [('' if v is None else str(v).replace('\n', ' / ').strip()) for v in row]
    if vals[0]:
        module = vals[0]
    text = ' || '.join(f'{chr(65+j)}={v}' for j, v in enumerate(vals) if v)
    if text:
        print(f'{i:03d} [{module}] {text}')
