import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')
p = r'd:\Users\HP\Downloads\xuanshiai problem\用户端测试问题.xlsx'
wb = openpyxl.load_workbook(p, data_only=True)
print('sheets:', [(ws.title, ws.max_row, ws.max_column) for ws in wb.worksheets])
for ws in wb.worksheets:
    print('=' * 20, ws.title)
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True), start=1):
        vals = [('' if v is None else str(v).replace('\n', ' / ').strip()) for v in row]
        if any(vals):
            print(f'{i:03d} || ' + ' | '.join(f'{chr(65+j)}={v}' for j, v in enumerate(vals) if v))
