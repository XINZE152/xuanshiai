"""单元测试：跟进全览 + 历史跟进导入（M3-5）管理后台能力。

覆盖范围：
- ``MemberFollowUpSummary`` / ``MemberFollowUpRow`` / ``MemberFollowUpListPage`` /
  ``MemberFollowUpImportResult`` 的可构造性与默认值；
- 新增路由已在 OpenAPI 注册（summary / import-template / import）；
- 未登录访问读 / 写接口均返回 401；
- 日期区间参数格式校验（422）；
- 导入模板可生成且能被重新解析（表头与列定位一致）；
- 导入解析辅助函数（时间格式 / 跟进方式 / 空值归一化）行为正确。
"""

import json
from datetime import datetime
from io import BytesIO

import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.main import app
from app.schemas.member_follow_up_admin import (
    MemberFollowUpImportResult,
    MemberFollowUpListPage,
    MemberFollowUpRow,
    MemberFollowUpSummary,
)
from app.services.member_follow_up_admin import (
    _cell_text,
    _find_column,
    _parse_datetime,
    _parse_method,
    _read_sheet,
    build_import_template,
)

client = TestClient(app)

BASE = "/api/v1/admin/members"


# ─── 路由注册（OpenAPI 路径断言） ───────────────────────────────


def test_follow_up_routes_are_registered() -> None:
    paths = client.get("/openapi.json").json()["paths"]
    assert "get" in paths[f"{BASE}/follow-ups"]
    assert "get" in paths[f"{BASE}/follow-ups/summary"]
    assert "get" in paths[f"{BASE}/follow-ups/import-template"]
    assert "post" in paths[f"{BASE}/follow-ups/import"]
    # 单会员跟进记录接口保持可用
    assert "get" in paths[f"{BASE}/{{member_id}}/follow-ups"]


def test_follow_up_routes_require_authentication() -> None:
    # 读接口需 matchmaker.member.read，未登录应 401
    assert client.get(f"{BASE}/follow-ups").status_code == 401
    assert client.get(f"{BASE}/follow-ups/summary").status_code == 401
    assert client.get(f"{BASE}/follow-ups/import-template").status_code == 401
    # 导入需 matchmaker.member.manage，未登录应 401
    assert client.post(f"{BASE}/follow-ups/import").status_code == 401


def test_follow_up_date_params_are_pattern_constrained() -> None:
    spec = client.get("/openapi.json").json()
    params = spec["paths"][f"{BASE}/follow-ups"]["get"]["parameters"]
    # Optional[...] 的约束会嵌在 anyOf 下，直接对参数 schema 做文本匹配
    start_schema = json.dumps(next(p for p in params if p["name"] == "start_date")["schema"])
    end_schema = json.dumps(next(p for p in params if p["name"] == "end_date")["schema"])
    assert "d{4}" in start_schema
    assert "d{4}" in end_schema
    intention_schema = json.dumps(next(p for p in params if p["name"] == "intention_level")["schema"])
    assert "minimum" in intention_schema
    assert "maximum" in intention_schema


# ─── 返回模型可构造 ─────────────────────────────────────────────


def test_follow_up_summary_defaults() -> None:
    summary = MemberFollowUpSummary()
    assert summary.all == 0
    assert summary.today == 0
    assert summary.yesterday == 0
    assert summary.three_days == 0
    assert summary.this_week == 0
    assert summary.last_week == 0
    assert summary.this_month == 0
    assert summary.last_month == 0


def test_follow_up_row_optional_fields_default_none() -> None:
    row = MemberFollowUpRow(
        id=1,
        user_id=10,
        member_code="G000010",
        matchmaker_name="admin",
        method="PHONE",
        method_label="电话",
        content="电话沟通",
    )
    assert row.nickname is None
    assert row.avatar is None
    assert row.note is None
    assert row.intention_level is None
    assert row.created_at is None


def test_follow_up_list_page_construct() -> None:
    page = MemberFollowUpListPage(
        items=[
            MemberFollowUpRow(
                id=1,
                user_id=10,
                member_code="G000010",
                matchmaker_name="芸希老师",
                method="WECHAT",
                method_label="微信",
                content="微信沟通",
            )
        ],
        page=1,
        page_size=20,
        total=1,
        has_more=False,
    )
    assert len(page.items) == 1
    assert page.items[0].member_code == "G000010"
    assert page.has_more is False


def test_follow_up_list_page_empty() -> None:
    page = MemberFollowUpListPage(items=[], page=1, page_size=20, total=0, has_more=False)
    assert page.items == []
    assert page.total == 0


def test_follow_up_import_result_defaults() -> None:
    result = MemberFollowUpImportResult()
    assert result.created == 0
    assert result.skipped == 0
    assert result.failed == 0
    assert result.errors == []


# ─── 导入模板 ───────────────────────────────────────────────────


def test_import_template_is_valid_xlsx() -> None:
    data = build_import_template()
    # xlsx 本质是 zip，魔数为 PK
    assert data[:2] == b"PK"
    workbook = load_workbook(BytesIO(data))
    sheet = workbook.active
    header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    assert list(header_row) == ["会员手机号", "跟进内容", "跟进人", "跟进时间", "跟进方式"]
    # 模板自带一行示例，便于管理员对齐格式
    sample_row = next(sheet.iter_rows(min_row=2, max_row=2, values_only=True))
    assert len(list(sample_row)) == 5
    workbook.close()


def test_import_template_can_be_parsed_back() -> None:
    headers, rows = _read_sheet(build_import_template(), "follow-up-import-template.xlsx")
    assert headers == ["会员手机号", "跟进内容", "跟进人", "跟进时间", "跟进方式"]
    # 模板含示例行
    assert len(rows) == 1
    excel_row, values = rows[0]
    assert excel_row == 2
    assert values[0] == "13800000000"


def test_read_sheet_rejects_other_formats() -> None:
    with pytest.raises(HTTPException) as excinfo:
        _read_sheet(b"PK\x03\x04", "follow-ups.xls")
    assert excinfo.value.status_code == 400
    assert ".xls" in excinfo.value.detail

    with pytest.raises(HTTPException) as excinfo:
        _read_sheet(b"PK\x03\x04", "follow-ups.csv")
    assert excinfo.value.status_code == 400


def test_read_sheet_rejects_template_missing_columns() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["姓名", "备注"])
    buffer = BytesIO()
    workbook.save(buffer)

    with pytest.raises(HTTPException) as excinfo:
        _read_sheet(buffer.getvalue(), "bad.xlsx")
    assert excinfo.value.status_code == 400


# ─── 解析辅助函数 ───────────────────────────────────────────────


def test_cell_text_normalizes_blank_values() -> None:
    assert _cell_text(None) == ""
    assert _cell_text("  ") == ""
    assert _cell_text("nan") == ""
    assert _cell_text("None") == ""
    assert _cell_text(" 张三 ") == "张三"
    assert _cell_text(13800000000) == "13800000000"
    assert _cell_text(3.0) == "3.0"


def test_parse_datetime_accepts_common_formats() -> None:
    assert _parse_datetime("2026-01-02 10:30:00") == datetime(2026, 1, 2, 10, 30, 0)
    assert _parse_datetime("2026-01-02 10:30") == datetime(2026, 1, 2, 10, 30)
    assert _parse_datetime("2026-01-02") == datetime(2026, 1, 2)
    assert _parse_datetime("2026/01/02 10:30:00") == datetime(2026, 1, 2, 10, 30, 0)
    native = datetime(2026, 3, 4, 12, 0, 0)
    assert _parse_datetime(native) is native


def test_parse_datetime_returns_none_for_unparsable() -> None:
    # 无法匹配时返回 None，由调用方回退为导入时间
    assert _parse_datetime("前天下午") is None
    assert _parse_datetime("") is None
    assert _parse_datetime(None) is None


def test_parse_method_maps_cn_and_en() -> None:
    assert _parse_method("电话") == "PHONE"
    assert _parse_method("微信") == "WECHAT"
    assert _parse_method("到店") == "VISIT"
    assert _parse_method("phone") == "PHONE"
    assert _parse_method("WeChat") == "WECHAT"
    # 无法识别 / 留空统一按 OTHER
    assert _parse_method("随便聊聊") == "OTHER"
    assert _parse_method("") == "OTHER"
    assert _parse_method(None) == "OTHER"


def test_find_column_matches_by_keyword() -> None:
    headers = ["会员手机号", "跟进内容", "跟进人", "跟进时间", "跟进方式"]
    assert _find_column(headers, "手机", "电话", "phone") == 0
    assert _find_column(headers, "跟进内容", "内容", "content") == 1
    assert _find_column(headers, "跟进人", "红娘", "matchmaker") == 2
    assert _find_column(headers, "跟进时间", "时间", "time") == 3
    assert _find_column(headers, "跟进方式", "方式", "method") == 4
    assert _find_column(headers, "不存在的列") is None
