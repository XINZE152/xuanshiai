"""Activity and signup management for the independent back office (M7)."""

from __future__ import annotations

import json
from io import BytesIO
from typing import Any

from fastapi import APIRouter, Depends, HTTPException, Path, Query, Response
from openpyxl import Workbook
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.dependencies import CurrentMatchmakerAdmin, get_current_matchmaker_admin
from app.db.session import get_db
from app.schemas.activity_admin import (
    ActivityAdminCreate,
    ActivityAdminItem,
    ActivityAdminPage,
    ActivityAdminUpdate,
    ActivityLinkInfo,
    ActivityOption,
    ActivitySignupAdminItem,
    ActivitySignupAdminPage,
    ActivitySignupStatistics,
    ActivitySignupUpdate,
    ActivityStatusUpdate,
)

router = APIRouter(prefix="/admin/activities")
signup_router = APIRouter(prefix="/admin/activity-signups")

_XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_ACTIVITY_BASE_URL = "https://www.xuanshi.com/subpages/active/index"

SELECT_ACTIVITY = "SELECT * FROM offline_activity"


def _bool(value: Any) -> bool:
    try:
        return bool(int(value))
    except (TypeError, ValueError):
        return bool(value)


def _num(value: Any) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def _snapshot_json(row: dict[str, Any]) -> str:
    """把审计前快照序列化为 JSON 文本（datetime/Decimal 统一转字符串）。"""
    return json.dumps(
        {
            key: (value.isoformat() if hasattr(value, "isoformat") else value)
            for key, value in row.items()
        },
        ensure_ascii=False,
        default=str,
    )


def _activity_item(row: Any) -> ActivityAdminItem:
    m = dict(row)
    return ActivityAdminItem(
        id=int(m["id"]),
        title=m["title"],
        cover=m.get("cover"),
        type=m.get("type"),
        city=m.get("city"),
        address=m.get("address"),
        start_time=m["start_time"],
        end_time=m["end_time"],
        signup_deadline=m.get("signup_deadline"),
        max_people=int(m.get("max_people") or 0),
        current_people=int(m.get("current_people") or 0),
        price=_num(m.get("price")),
        status=int(m.get("status") or 1),
        description=m.get("description"),
        created_by=m.get("created_by"),
        created_at=m["created_at"],
        organizer=m.get("organizer"),
        cover_small=m.get("cover_small"),
        fee_name=m.get("fee_name") or "报名费",
        price_male=_num(m.get("price_male")),
        price_female=_num(m.get("price_female")),
        signup_mode=m.get("signup_mode") or "anyone",
        require_realname=_bool(m.get("require_realname")),
        limit_mode=m.get("limit_mode") or "gender",
        max_male=int(m.get("max_male") or 0),
        max_female=int(m.get("max_female") or 0),
        virtual_people=int(m.get("virtual_people") or 0),
        virtual_female=int(m.get("virtual_female") or 0),
        hide_signup_count=_bool(m.get("hide_signup_count")),
        reward_promoter=_num(m.get("reward_promoter")),
        reward_service=_num(m.get("reward_service")),
        reward_partner=_num(m.get("reward_partner")),
        reminder_html=m.get("reminder_html"),
        service_wechat=m.get("service_wechat"),
        service_qr=m.get("service_qr"),
        virtual_views=int(m.get("virtual_views") or 0),
        sort_order=int(m.get("sort_order") or 0),
        custom_share=_bool(m.get("custom_share")),
        manager_ids=m.get("manager_ids"),
        notify_phones=m.get("notify_phones"),
        online=_bool(m.get("online") if m.get("online") is not None else 1),
        audit_status=m.get("audit_status") or "approved",
        male_count=int(m.get("male_count") or 0),
        female_count=int(m.get("female_count") or 0),
        link_url=f"{_ACTIVITY_BASE_URL}?id={int(m['id'])}",
    )


def _signup_item(row: Any) -> ActivitySignupAdminItem:
    m = dict(row)
    return ActivitySignupAdminItem(
        id=int(m["id"]),
        activity_id=int(m["activity_id"]),
        activity_title=m.get("activity_title"),
        user_id=int(m["user_id"]),
        nickname=m.get("nickname"),
        real_name=m.get("real_name"),
        phone=m.get("phone"),
        remark=m.get("remark"),
        status=int(m.get("status") or 0),
        cancel_reason=m.get("cancel_reason"),
        created_at=m["created_at"],
        updated_at=m["updated_at"],
        gender=m.get("gender"),
        age=m.get("age"),
        height=m.get("height"),
        education=m.get("education"),
        income=m.get("income"),
        marriage_status=m.get("marriage_status"),
        company=m.get("company"),
        avatar=m.get("avatar") or m.get("user_avatar"),
        id_card=m.get("id_card"),
        is_member=_bool(m.get("is_member")),
        is_realname=_bool(m.get("is_realname")),
        signup_times=int(m.get("signup_times") or 1),
        pay_status=m.get("pay_status") or "free",
        pay_amount=_num(m.get("pay_amount")),
        checked_in=_bool(m.get("checked_in")),
        in_crm=_bool(m.get("in_crm")),
        promoter_id=m.get("promoter_id"),
        promoter_name=m.get("promoter_name"),
    )


async def _get_activity(db: AsyncSession, activity_id: int) -> ActivityAdminItem:
    row = (
        await db.execute(
            text(
                "SELECT a.*, "
                "(SELECT COUNT(*) FROM activity_signup s JOIN users u ON u.id = s.user_id "
                " WHERE s.activity_id = a.id AND s.status = 1 AND s.deleted_at IS NULL "
                " AND COALESCE(s.gender, CASE u.gender WHEN 1 THEN '男' WHEN 2 THEN '女' END) = '男') AS male_count, "
                "(SELECT COUNT(*) FROM activity_signup s JOIN users u ON u.id = s.user_id "
                " WHERE s.activity_id = a.id AND s.status = 1 AND s.deleted_at IS NULL "
                " AND COALESCE(s.gender, CASE u.gender WHEN 1 THEN '男' WHEN 2 THEN '女' END) = '女') AS female_count "
                "FROM offline_activity a WHERE a.id = :id"
            ),
            {"id": activity_id},
        )
    ).mappings().first()
    if not row:
        raise HTTPException(404, detail="活动不存在")
    return _activity_item(row)


@router.get("", response_model=ActivityAdminPage, summary="查询活动列表")
async def activities(
    page: int = Query(1, ge=1, le=1000),
    page_size: int = Query(20, ge=1, le=100),
    status: int | None = Query(None, ge=1, le=5),
    city: str | None = Query(None, max_length=64),
    search: str | None = Query(None, max_length=128),
    online: bool | None = Query(None),
    audit_status: str | None = Query(None, pattern="^(pending|approved|rejected)$"),
    current: CurrentMatchmakerAdmin = Depends(get_current_matchmaker_admin),
    db: AsyncSession = Depends(get_db),
) -> ActivityAdminPage:
    current.require("community.activity.read")
    where = ["1=1"]
    params: dict = {"limit": page_size, "offset": (page - 1) * page_size}
    if status is not None:
        where.append("status = :status")
        params["status"] = status
    if city:
        where.append("city = :city")
        params["city"] = city
    if search:
        where.append("title LIKE CONCAT('%', :search, '%')")
        params["search"] = search
    if online is not None:
        where.append("COALESCE(online, 1) = :online")
        params["online"] = 1 if online else 0
    if audit_status:
        where.append("COALESCE(audit_status, 'approved') = :audit_status")
        params["audit_status"] = audit_status
    clause = " AND ".join(where)
    rows = await db.execute(
        text(
            "SELECT a.*, "
            "(SELECT COUNT(*) FROM activity_signup s JOIN users u ON u.id = s.user_id "
            " WHERE s.activity_id = a.id AND s.status = 1 AND s.deleted_at IS NULL "
            " AND COALESCE(s.gender, CASE u.gender WHEN 1 THEN '男' WHEN 2 THEN '女' END) = '男') AS male_count, "
            "(SELECT COUNT(*) FROM activity_signup s JOIN users u ON u.id = s.user_id "
            " WHERE s.activity_id = a.id AND s.status = 1 AND s.deleted_at IS NULL "
            " AND COALESCE(s.gender, CASE u.gender WHEN 1 THEN '男' WHEN 2 THEN '女' END) = '女') AS female_count "
            f"FROM offline_activity a WHERE {clause} ORDER BY a.id DESC LIMIT :limit OFFSET :offset"
        ),
        params,
    )
    count = await db.execute(
        text(f"SELECT COUNT(*) FROM offline_activity a WHERE {clause}"),
        {k: v for k, v in params.items() if k not in ("limit", "offset")},
    )
    total = int(count.scalar() or 0)
    return ActivityAdminPage(
        items=[_activity_item(row) for row in rows.mappings().all()],
        page=page,
        page_size=page_size,
        total=total,
        has_more=page * page_size < total,
    )


@router.post("", response_model=ActivityAdminItem, status_code=201, summary="创建活动")
async def create_activity(
    body: ActivityAdminCreate,
    current: CurrentMatchmakerAdmin = Depends(get_current_matchmaker_admin),
    db: AsyncSession = Depends(get_db),
) -> ActivityAdminItem:
    current.require("community.activity.manage")
    payload = body.model_dump()
    payload = {k: (int(v) if isinstance(v, bool) else v) for k, v in payload.items()}
    payload["created_by"] = current.account.id
    columns = ", ".join(payload.keys())
    placeholders = ", ".join(f":{key}" for key in payload)
    result = await db.execute(
        text(f"INSERT INTO offline_activity ({columns}) VALUES ({placeholders})"),
        payload,
    )
    await db.commit()
    return await _get_activity(db, int(result.lastrowid))


@router.get("/options", response_model=list[ActivityOption], summary="活动下拉字典")
async def activity_options(
    current: CurrentMatchmakerAdmin = Depends(get_current_matchmaker_admin),
    db: AsyncSession = Depends(get_db),
) -> list[ActivityOption]:
    current.require("community.activity.read")
    rows = await db.execute(text("SELECT id, title FROM offline_activity ORDER BY id DESC LIMIT 500"))
    return [ActivityOption(id=int(r[0]), title=str(r[1])) for r in rows.all()]


@router.get("/{activity_id}/signups", response_model=ActivitySignupAdminPage, summary="查询活动报名")
async def signups(
    activity_id: int = Path(..., ge=1),
    page: int = Query(1, ge=1, le=1000),
    page_size: int = Query(20, ge=1, le=100),
    status: int | None = Query(None, ge=0, le=3),
    current: CurrentMatchmakerAdmin = Depends(get_current_matchmaker_admin),
    db: AsyncSession = Depends(get_db),
) -> ActivitySignupAdminPage:
    current.require("community.activity.read")
    await _get_activity(db, activity_id)
    where = ["s.activity_id = :activity_id", "s.deleted_at IS NULL"]
    params: dict = {"activity_id": activity_id, "limit": page_size, "offset": (page - 1) * page_size}
    if status is not None:
        where.append("s.status = :status")
        params["status"] = status
    clause = " AND ".join(where)
    rows = await db.execute(
        text(f"SELECT s.*, u.nickname, u.avatar AS user_avatar, a.title AS activity_title, p.nickname AS promoter_name "
             f"FROM activity_signup s LEFT JOIN users u ON u.id = s.user_id "
             f"LEFT JOIN offline_activity a ON a.id = s.activity_id "
             f"LEFT JOIN users p ON p.id = s.promoter_id "
             f"WHERE {clause} ORDER BY s.id DESC LIMIT :limit OFFSET :offset"),
        params,
    )
    count = await db.execute(
        text(f"SELECT COUNT(*) FROM activity_signup s WHERE {clause}"),
        {k: v for k, v in params.items() if k not in ("limit", "offset")},
    )
    total = int(count.scalar() or 0)
    return ActivitySignupAdminPage(
        items=[_signup_item(row) for row in rows.mappings().all()],
        page=page,
        page_size=page_size,
        total=total,
        has_more=page * page_size < total,
    )


@router.get("/{activity_id}/link", response_model=ActivityLinkInfo, summary="活动链接与二维码")
async def activity_link(
    activity_id: int = Path(..., ge=1),
    current: CurrentMatchmakerAdmin = Depends(get_current_matchmaker_admin),
) -> ActivityLinkInfo:
    current.require("community.activity.read")
    return ActivityLinkInfo(link_url=f"{_ACTIVITY_BASE_URL}?id={activity_id}", qr_code=None)


@router.post("/{activity_id}/copy", response_model=ActivityAdminItem, status_code=201, summary="复制活动")
async def copy_activity(
    activity_id: int = Path(..., ge=1),
    current: CurrentMatchmakerAdmin = Depends(get_current_matchmaker_admin),
    db: AsyncSession = Depends(get_db),
) -> ActivityAdminItem:
    current.require("community.activity.manage")
    await _get_activity(db, activity_id)
    await db.execute(
        text(
            "INSERT INTO offline_activity "
            "(title, cover, cover_small, type, city, address, start_time, end_time, signup_deadline, "
            " max_people, price, description, organizer, time_text, fee_name, price_male, price_female, signup_mode, "
            " require_realname, limit_mode, max_male, max_female, virtual_people, virtual_female, "
            " hide_signup_count, reward_promoter, reward_service, reward_partner, reminder_html, "
            " service_wechat, service_qr, virtual_views, sort_order, custom_share, manager_ids, "
            " notify_phones, status, online, audit_status, created_by) "
            "SELECT CONCAT(title, '(副本)'), cover, cover_small, type, city, address, start_time, end_time, "
            " signup_deadline, max_people, price, description, organizer, time_text, fee_name, price_male, price_female, "
            " signup_mode, require_realname, limit_mode, max_male, max_female, virtual_people, virtual_female, "
            " hide_signup_count, reward_promoter, reward_service, reward_partner, reminder_html, "
            " service_wechat, service_qr, virtual_views, sort_order, custom_share, manager_ids, "
            " notify_phones, 1, 0, 'pending', :actor "
            "FROM offline_activity WHERE id = :id"
        ),
        {"id": activity_id, "actor": current.account.id},
    )
    await db.commit()
    fresh = (await db.execute(text("SELECT MAX(id) AS id FROM offline_activity"))).scalar()
    return await _get_activity(db, int(fresh))


@router.patch("/{activity_id}/status", response_model=ActivityAdminItem, summary="修改活动状态")
async def update_activity_status(
    activity_id: int = Path(..., ge=1),
    body: ActivityStatusUpdate = ...,
    current: CurrentMatchmakerAdmin = Depends(get_current_matchmaker_admin),
    db: AsyncSession = Depends(get_db),
) -> ActivityAdminItem:
    current.require("community.activity.manage")
    await _get_activity(db, activity_id)
    await db.execute(
        text("UPDATE offline_activity SET status = :status, updated_at = UTC_TIMESTAMP() WHERE id = :id"),
        {"status": body.status, "id": activity_id},
    )
    await db.commit()
    return await _get_activity(db, activity_id)


@router.patch("/{activity_id}", response_model=ActivityAdminItem, summary="修改活动")
async def update_activity(
    activity_id: int = Path(..., ge=1),
    body: ActivityAdminUpdate = ...,
    current: CurrentMatchmakerAdmin = Depends(get_current_matchmaker_admin),
    db: AsyncSession = Depends(get_db),
) -> ActivityAdminItem:
    current.require("community.activity.manage")
    await _get_activity(db, activity_id)
    values = {k: (int(v) if isinstance(v, bool) else v) for k, v in body.model_dump(exclude_unset=True).items()}
    updates = ", ".join(f"{key} = :{key}" for key in values)
    await db.execute(
        text(f"UPDATE offline_activity SET {updates}, updated_at = UTC_TIMESTAMP() WHERE id = :id"),
        {**values, "id": activity_id},
    )
    await db.commit()
    return await _get_activity(db, activity_id)


@router.delete("/{activity_id}", status_code=204, summary="删除活动")
async def delete_activity(
    activity_id: int = Path(..., ge=1),
    current: CurrentMatchmakerAdmin = Depends(get_current_matchmaker_admin),
    db: AsyncSession = Depends(get_db),
) -> None:
    current.require("community.activity.manage")
    await _get_activity(db, activity_id)
    # 活动级联删除会把报名一起物理删除；已支付/已签到的报名涉及金额与履约记录，
    # 必须先走退款/撤销流程，避免直接丢失对账依据。
    blocking = (
        await db.execute(
            text("SELECT COUNT(*) FROM activity_signup WHERE activity_id = :id "
                 "AND deleted_at IS NULL AND (pay_status = 'paid' OR COALESCE(checked_in, 0) = 1)"),
            {"id": activity_id},
        )
    ).scalar()
    if int(blocking or 0) > 0:
        raise HTTPException(
            409,
            detail=f"该活动存在 {int(blocking)} 条已支付或已签到的报名，不能删除，请先处理退款或撤销签到",
        )
    await db.execute(text("DELETE FROM activity_signup WHERE activity_id = :id"), {"id": activity_id})
    await db.execute(text("DELETE FROM offline_activity WHERE id = :id"), {"id": activity_id})
    await db.execute(
        text("INSERT INTO business_audit_log (actor_user_id, action, resource_type, resource_id) "
             "VALUES (:actor, 'activity.delete', 'offline_activity', :id)"),
        {"actor": current.account.id, "id": activity_id},
    )
    await db.commit()


@router.get("/{activity_id}", response_model=ActivityAdminItem, summary="查询活动详情")
async def activity_detail(
    activity_id: int = Path(..., ge=1),
    current: CurrentMatchmakerAdmin = Depends(get_current_matchmaker_admin),
    db: AsyncSession = Depends(get_db),
) -> ActivityAdminItem:
    current.require("community.activity.read")
    return await _get_activity(db, activity_id)


# ─────────────────────────── 报名管理（独立前缀） ───────────────────────────

@signup_router.get("", response_model=ActivitySignupAdminPage, summary="分页查询活动报名")
async def signup_list(
    activity_id: int | None = Query(None, ge=1),
    page: int = Query(1, ge=1, le=1000),
    page_size: int = Query(20, ge=1, le=100),
    status: int | None = Query(None, ge=0, le=3),
    first_signup: bool | None = Query(None),
    gender: str | None = Query(None, max_length=8),
    pay_status: str | None = Query(None, pattern="^(free|paid|unpaid)$"),
    checked_in: bool | None = Query(None),
    keyword: str | None = Query(None, max_length=128),
    search_by: str | None = Query(None, pattern="^(nickname|phone)$"),
    signup_from: str | None = Query(None, pattern=r"^\d{4}-\d{2}-\d{2}$"),
    signup_to: str | None = Query(None, pattern=r"^\d{4}-\d{2}-\d{2}$"),
    current: CurrentMatchmakerAdmin = Depends(get_current_matchmaker_admin),
    db: AsyncSession = Depends(get_db),
) -> ActivitySignupAdminPage:
    """Return the same signup shape as the activity-scoped endpoint, including all activities."""
    current.require("community.activity.read")
    where = ["1=1", "s.deleted_at IS NULL"]
    params: dict[str, Any] = {
        "limit": page_size,
        "offset": (page - 1) * page_size,
    }
    if activity_id is not None:
        where.append("s.activity_id = :activity_id")
        params["activity_id"] = activity_id
    if status is not None:
        where.append("s.status = :status")
        params["status"] = status
    if first_signup is not None:
        where.append("COALESCE(s.signup_times, 1) <= 1" if first_signup else "COALESCE(s.signup_times, 1) > 1")
    if gender:
        where.append("s.gender = :gender")
        params["gender"] = gender
    if pay_status:
        where.append("s.pay_status = :pay_status")
        params["pay_status"] = pay_status
    if checked_in is not None:
        where.append("COALESCE(s.checked_in, 0) = :checked_in")
        params["checked_in"] = 1 if checked_in else 0
    if keyword:
        field = "u.phone" if search_by == "phone" else "COALESCE(u.nickname, s.real_name)"
        where.append(f"{field} LIKE CONCAT('%', :keyword, '%')")
        params["keyword"] = keyword
    if signup_from:
        where.append("s.created_at >= STR_TO_DATE(:signup_from, '%Y-%m-%d')")
        params["signup_from"] = signup_from
    if signup_to:
        where.append(
            "s.created_at < DATE_ADD(STR_TO_DATE(:signup_to, '%Y-%m-%d'), INTERVAL 1 DAY)"
        )
        params["signup_to"] = signup_to

    clause = " AND ".join(where)
    select_sql = (
        "SELECT s.*, u.nickname, u.avatar AS user_avatar, "
        "a.title AS activity_title, p.nickname AS promoter_name "
        "FROM activity_signup s "
        "LEFT JOIN users u ON u.id = s.user_id "
        "LEFT JOIN offline_activity a ON a.id = s.activity_id "
        "LEFT JOIN users p ON p.id = s.promoter_id "
        f"WHERE {clause} ORDER BY s.id DESC LIMIT :limit OFFSET :offset"
    )
    rows = await db.execute(text(select_sql), params)
    count = await db.execute(
        text(
            "SELECT COUNT(*) FROM activity_signup s "
            "LEFT JOIN users u ON u.id = s.user_id "
            f"WHERE {clause}"
        ),
        {key: value for key, value in params.items() if key not in ("limit", "offset")},
    )
    total = int(count.scalar() or 0)
    return ActivitySignupAdminPage(
        items=[_signup_item(row) for row in rows.mappings().all()],
        page=page,
        page_size=page_size,
        total=total,
        has_more=page * page_size < total,
    )


@signup_router.get("/statistics", response_model=ActivitySignupStatistics, summary="活动报名统计卡")
async def signup_statistics(
    activity_id: int | None = Query(None, ge=1),
    current: CurrentMatchmakerAdmin = Depends(get_current_matchmaker_admin),
    db: AsyncSession = Depends(get_db),
) -> ActivitySignupStatistics:
    current.require("community.activity.read")
    where = "WHERE 1=1 AND deleted_at IS NULL"
    params: dict = {}
    if activity_id:
        where += " AND activity_id = :activity_id"
        params["activity_id"] = activity_id
    row = (
        await db.execute(
            text(
                "SELECT COUNT(*) AS total, "
                "SUM(CASE WHEN signup_times <= 1 THEN 1 ELSE 0 END) AS first_signup, "
                "SUM(CASE WHEN status = 0 THEN 1 ELSE 0 END) AS pending, "
                "SUM(CASE WHEN status = 1 THEN 1 ELSE 0 END) AS approved, "
                "SUM(CASE WHEN status = 3 THEN 1 ELSE 0 END) AS rejected, "
                "COALESCE(SUM(pay_amount), 0) AS fee_amount, "
                "SUM(CASE WHEN COALESCE(checked_in,0) = 0 THEN 1 ELSE 0 END) AS not_checked_in, "
                "SUM(CASE WHEN COALESCE(checked_in,0) = 1 THEN 1 ELSE 0 END) AS checked_in, "
                "SUM(CASE WHEN COALESCE(in_crm,0) = 1 THEN 1 ELSE 0 END) AS in_crm "
                f"FROM activity_signup {where}"
            ),
            params,
        )
    ).mappings().first()
    m = dict(row or {})
    return ActivitySignupStatistics(
        total=int(m.get("total") or 0),
        first_signup=int(m.get("first_signup") or 0),
        pending=int(m.get("pending") or 0),
        approved=int(m.get("approved") or 0),
        rejected=int(m.get("rejected") or 0),
        fee_amount=format(_num(m.get("fee_amount")), ".2f"),
        not_checked_in=int(m.get("not_checked_in") or 0),
        checked_in=int(m.get("checked_in") or 0),
        in_crm=int(m.get("in_crm") or 0),
    )


@signup_router.get("/options", response_model=list[ActivityOption], summary="报名筛选活动下拉")
async def signup_options(
    current: CurrentMatchmakerAdmin = Depends(get_current_matchmaker_admin),
    db: AsyncSession = Depends(get_db),
) -> list[ActivityOption]:
    current.require("community.activity.read")
    rows = await db.execute(text("SELECT id, title FROM offline_activity ORDER BY id DESC LIMIT 500"))
    return [ActivityOption(id=int(r[0]), title=str(r[1])) for r in rows.all()]


@signup_router.get("/export", summary="导出活动报名 Excel")
async def signup_export(
    activity_id: int | None = Query(None, ge=1),
    status: int | None = Query(None, ge=0, le=3),
    current: CurrentMatchmakerAdmin = Depends(get_current_matchmaker_admin),
    db: AsyncSession = Depends(get_db),
) -> Response:
    current.require("community.activity.read")
    where = ["1=1", "s.deleted_at IS NULL"]
    params: dict = {}
    if activity_id:
        where.append("s.activity_id = :activity_id")
        params["activity_id"] = activity_id
    if status is not None:
        where.append("s.status = :status")
        params["status"] = status
    clause = " AND ".join(where)
    rows = (
        await db.execute(
            text(
                "SELECT s.id, a.title AS activity_title, s.user_id, u.nickname, s.real_name, s.gender, "
                "s.age, s.phone, s.signup_times, s.pay_status, s.pay_amount, s.status, s.created_at "
                "FROM activity_signup s LEFT JOIN users u ON u.id = s.user_id "
                "LEFT JOIN offline_activity a ON a.id = s.activity_id "
                f"WHERE {clause} ORDER BY s.id DESC"
            ),
            params,
        )
    ).mappings().all()
    status_label = {0: "待审", 1: "审核通过", 2: "已取消", 3: "未通过"}
    wb = Workbook()
    ws = wb.active
    ws.title = "活动报名"
    ws.append(["ID", "活动", "会员ID", "昵称", "姓名", "性别", "年龄", "手机", "第几次报名", "缴费状态", "缴费金额", "审核状态", "报名时间"])
    for r in rows:
        ws.append([
            int(r["id"]), r["activity_title"], int(r["user_id"]), r["nickname"], r["real_name"],
            r["gender"], r["age"], r["phone"], int(r["signup_times"] or 1), r["pay_status"],
            _num(r["pay_amount"]), status_label.get(int(r["status"] or 0), "待审"),
            r["created_at"].strftime("%Y-%m-%d %H:%M:%S") if r["created_at"] else "",
        ])
    buffer = BytesIO()
    wb.save(buffer)
    return Response(
        content=buffer.getvalue(),
        media_type=_XLSX_MEDIA_TYPE,
        headers={"Content-Disposition": 'attachment; filename="activity-signups.xlsx"'},
    )


@signup_router.get("/{signup_id}", response_model=ActivitySignupAdminItem, summary="查询活动报名详情")
async def signup_detail(
    signup_id: int = Path(..., ge=1),
    current: CurrentMatchmakerAdmin = Depends(get_current_matchmaker_admin),
    db: AsyncSession = Depends(get_db),
) -> ActivitySignupAdminItem:
    current.require("community.activity.read")
    row = (
        await db.execute(
            text("SELECT s.*, u.nickname, u.avatar AS user_avatar, a.title AS activity_title, p.nickname AS promoter_name "
                 "FROM activity_signup s LEFT JOIN users u ON u.id = s.user_id "
                 "LEFT JOIN offline_activity a ON a.id = s.activity_id "
                 "LEFT JOIN users p ON p.id = s.promoter_id WHERE s.id = :id AND s.deleted_at IS NULL"),
            {"id": signup_id},
        )
    ).mappings().first()
    if not row:
        raise HTTPException(404, detail="活动报名不存在")
    return _signup_item(row)


@signup_router.patch("/{signup_id}", response_model=ActivitySignupAdminItem, summary="修改报名/审核报名")
async def update_signup(
    signup_id: int = Path(..., ge=1),
    body: ActivitySignupUpdate = ...,
    current: CurrentMatchmakerAdmin = Depends(get_current_matchmaker_admin),
    db: AsyncSession = Depends(get_db),
) -> ActivitySignupAdminItem:
    current.require("community.activity.manage")
    values = {k: (int(v) if isinstance(v, bool) else v) for k, v in body.model_dump(exclude_unset=True).items()}
    if not values:
        raise HTTPException(422, detail="至少提供一个需要修改的字段")
    updates = ", ".join(f"{key} = :{key}" for key in values)
    result = await db.execute(
        text(f"UPDATE activity_signup SET {updates}, updated_at = UTC_TIMESTAMP() "
             "WHERE id = :id AND deleted_at IS NULL"),
        {**values, "id": signup_id},
    )
    if result.rowcount == 0:
        raise HTTPException(404, detail="活动报名不存在")
    await db.execute(
        text("INSERT INTO business_audit_log (actor_user_id, action, resource_type, resource_id) "
             "VALUES (:actor, 'activity_signup.update', 'activity_signup', :id)"),
        {"actor": current.account.id, "id": signup_id},
    )
    await db.commit()
    row = (
        await db.execute(
            text("SELECT s.*, u.nickname, u.avatar AS user_avatar, a.title AS activity_title, p.nickname AS promoter_name "
                 "FROM activity_signup s LEFT JOIN users u ON u.id = s.user_id "
                 "LEFT JOIN offline_activity a ON a.id = s.activity_id "
                 "LEFT JOIN users p ON p.id = s.promoter_id WHERE s.id = :id"),
            {"id": signup_id},
        )
    ).mappings().one()
    return _signup_item(row)


@signup_router.delete("/{signup_id}", status_code=204, summary="删除活动报名（软删除并留审计）")
async def delete_signup(
    signup_id: int = Path(..., ge=1),
    reason: str | None = Query(None, max_length=255),
    current: CurrentMatchmakerAdmin = Depends(get_current_matchmaker_admin),
    db: AsyncSession = Depends(get_db),
) -> None:
    """软删除报名记录。

    报名记录含支付金额与签到状态，物理删除后无法追溯、也无法支撑退款对账，
    因此改为软删除并写入 `business_audit_log`（含删除前快照）。已支付或已签到的
    报名禁止删除，需走退款/撤销流程。
    """
    current.require("community.activity.manage")
    row = (
        await db.execute(
            text("SELECT * FROM activity_signup WHERE id = :id AND deleted_at IS NULL FOR UPDATE"),
            {"id": signup_id},
        )
    ).mappings().first()
    if not row:
        raise HTTPException(404, detail="活动报名不存在")
    snapshot = dict(row)
    if str(snapshot.get("pay_status") or "") == "paid":
        raise HTTPException(409, detail="该报名已支付，不能删除，请走退款流程")
    if int(snapshot.get("checked_in") or 0) == 1:
        raise HTTPException(409, detail="该报名已签到，不能删除，请先撤销签到")
    await db.execute(
        text("UPDATE activity_signup SET deleted_at = UTC_TIMESTAMP(), deleted_by = :actor, "
             "updated_at = UTC_TIMESTAMP() WHERE id = :id AND deleted_at IS NULL"),
        {"actor": current.account.id, "id": signup_id},
    )
    # 名额回收：被删报名若处于「待审核 / 已通过」，此前已计入 current_people，必须同事务减回去，
    # 否则用户端剩余名额与「名额已满」状态会长期偏小。SET 按从左到右求值，故 status 判断用的是已减后的值。
    if int(snapshot.get("status") or 0) in (0, 1):
        await db.execute(
            text("UPDATE offline_activity SET "
                 "current_people = CASE WHEN current_people > 0 THEN current_people - 1 ELSE 0 END, "
                 "status = CASE WHEN status = 2 AND (max_people <= 0 OR current_people < max_people) "
                 "THEN 1 ELSE status END, updated_at = UTC_TIMESTAMP() WHERE id = :activity_id"),
            {"activity_id": snapshot["activity_id"]},
        )
    await db.execute(
        text("INSERT INTO business_audit_log "
             "(actor_user_id, action, resource_type, resource_id, before_json, reason) "
             "VALUES (:actor, 'activity_signup.delete', 'activity_signup', :id, :before_json, :reason)"),
        {
            "actor": current.account.id,
            "id": signup_id,
            "before_json": _snapshot_json(snapshot),
            "reason": reason or "后台删除报名",
        },
    )
    await db.commit()
