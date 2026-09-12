"""Customer lead management for the independent back office."""

from datetime import datetime
from io import BytesIO
from typing import Any

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from sqlalchemy import text
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession

from app.schemas.customer_lead_admin import (
    CustomerLead, CustomerLeadAbandonment, CustomerLeadAssignment, CustomerLeadCreate, CustomerLeadFollowUp,
    CustomerLeadBatchImportResult, CustomerLeadFollowUpCreate, CustomerLeadImportRow, CustomerLeadImportSummary,
    CustomerLeadOption, CustomerLeadOptions,
    CustomerLeadPage, CustomerLeadStatistics, CustomerLeadUpdate,
)
from app.services.customer_lead_contact import (
    ensure_contact_available, normalize_contact, raise_duplicate_contact,
)

LEAD_SELECT = """SELECT id, name, phone, wechat, source, intention_level, status, audit_status, matchmaker_id,
    organization_id, promoter_id, next_follow_at, converted_user_id, remark, created_by, created_at, updated_at,
    (SELECT GROUP_CONCAT(t.name ORDER BY t.id SEPARATOR ',')
        FROM customer_lead_tag_relation r JOIN customer_lead_tag t ON t.id = r.tag_id
        WHERE r.lead_id = customer_lead.id) AS tags_raw
    FROM customer_lead"""

# 导入模板表头（顺序即模板列顺序）
_TEMPLATE_HEADERS: list[str] = ["姓名", "手机号", "微信号", "来源", "意向程度", "备注"]
_TEMPLATE_SAMPLE: list[str] = ["示例：张女士", "13800000000", "zhang_wx", "抖音", "中", "示例行，导入前请删除"]
_INTENTION_ALIASES: dict[str, int] = {"低": 1, "中": 2, "高": 3, "1": 1, "2": 2, "3": 3}


def _lead(row: Any) -> CustomerLead:
    payload = dict(row)
    raw = payload.pop("tags_raw", None)
    payload["tags"] = [item for item in (raw or "").split(",") if item]
    return CustomerLead(**payload)


async def _validate_owner(db: AsyncSession, assignment: CustomerLeadAssignment) -> None:
    if assignment.matchmaker_id is not None:
        result = await db.execute(text("""SELECT 1 FROM user_matchmaker_apply a JOIN user_role r ON r.user_id = a.user_id
            AND r.role_code = 'service_matchmaker' AND r.status = 1
            WHERE a.user_id = :id AND a.application_type = 'service_matchmaker' AND a.status = 1"""), {"id": assignment.matchmaker_id})
        if not result.scalar():
            raise HTTPException(422, detail="只能分配给有效服务红娘")
    if assignment.organization_id is not None:
        result = await db.execute(text("SELECT 1 FROM organization WHERE id = :id AND org_type = 'store' AND status = 1"), {"id": assignment.organization_id})
        if not result.scalar():
            raise HTTPException(422, detail="门店不存在或已停用")


async def _sync_tags(db: AsyncSession, lead_id: int, tags: list[str]) -> None:
    """按名称写入客源标签关联；不存在的标签名自动补建配置行。"""
    names = [name.strip() for name in tags if name and name.strip()]
    if not names:
        return
    await db.execute(text("DELETE FROM customer_lead_tag_relation WHERE lead_id = :id"), {"id": lead_id})
    for name in dict.fromkeys(names):
        await db.execute(text("""INSERT IGNORE INTO customer_lead_tag (name, enabled, sort) VALUES (:name, 1, 0)"""), {"name": name})
        tag_id = (await db.execute(text("SELECT id FROM customer_lead_tag WHERE name = :name"), {"name": name})).scalar()
        if tag_id:
            await db.execute(text("""INSERT IGNORE INTO customer_lead_tag_relation (lead_id, tag_id) VALUES (:lead_id, :tag_id)"""), {"lead_id": lead_id, "tag_id": int(tag_id)})


async def create_lead(db: AsyncSession, account_id: int, request: CustomerLeadCreate) -> CustomerLead:
    phone = normalize_contact(request.phone)
    wechat = normalize_contact(request.wechat)
    await ensure_contact_available(db, phone, wechat)
    try:
        result = await db.execute(text("""INSERT INTO customer_lead
            (name, phone, wechat, source, intention_level, remark, promoter_id, audit_status, created_by)
            VALUES (:name, :phone, :wechat, :source, :intention_level, :remark, :promoter_id, :audit_status, :created_by)"""), {
            **request.model_dump(exclude={"tags"}), "phone": phone, "wechat": wechat, "created_by": account_id,
        })
        lead_id = int(result.lastrowid)
        await _sync_tags(db, lead_id, request.tags)
        await db.execute(text("""INSERT INTO business_audit_log (actor_user_id, action, resource_type, resource_id)
            VALUES (:actor, 'customer_lead.create', 'customer_lead', :id)"""), {"actor": account_id, "id": lead_id})
        await db.commit()
    except IntegrityError:
        await db.rollback()
        raise_duplicate_contact()
    return await get_lead(db, lead_id)


async def get_lead(db: AsyncSession, lead_id: int) -> CustomerLead:
    row = (await db.execute(text(f"{LEAD_SELECT} WHERE id = :id"), {"id": lead_id})).mappings().first()
    if not row:
        raise HTTPException(404, detail="客源线索不存在")
    return _lead(row)


async def list_leads(db: AsyncSession, page: int, page_size: int, status: str | None, source: str | None, matchmaker_id: int | None, search: str | None) -> CustomerLeadPage:
    where = ["1=1"]
    params: dict[str, Any] = {"limit": page_size, "offset": (page - 1) * page_size}
    if status:
        where.append("status = :status")
        params["status"] = status
    if source:
        where.append("source = :source")
        params["source"] = source
    if matchmaker_id is not None:
        where.append("matchmaker_id = :matchmaker_id")
        params["matchmaker_id"] = matchmaker_id
    if search:
        where.append("(name LIKE CONCAT('%', :search, '%') OR phone LIKE CONCAT('%', :search, '%') OR wechat LIKE CONCAT('%', :search, '%'))")
        params["search"] = search
    clause = " AND ".join(where)
    rows = await db.execute(text(f"{LEAD_SELECT} WHERE {clause} ORDER BY id DESC LIMIT :limit OFFSET :offset"), params)
    count = await db.execute(text(f"SELECT COUNT(*) FROM customer_lead WHERE {clause}"), {k: v for k, v in params.items() if k not in ("limit", "offset")})
    total = int(count.scalar() or 0)
    return CustomerLeadPage(items=[_lead(row) for row in rows.mappings().all()], page=page, page_size=page_size, total=total, has_more=page * page_size < total)


async def update_lead(db: AsyncSession, account_id: int, lead_id: int, request: CustomerLeadUpdate) -> CustomerLead:
    current = await get_lead(db, lead_id)
    values = request.model_dump(exclude_unset=True)
    target_status = values.get("status") or current.status
    # 目标状态仍为有效线索时，联系方式不得与其它有效线索重复；改为 LOST/CLOSED 后生成列会置空。
    if target_status not in ("LOST", "CLOSED"):
        target_phone = normalize_contact(values["phone"]) if "phone" in values else normalize_contact(current.phone)
        target_wechat = normalize_contact(values["wechat"]) if "wechat" in values else normalize_contact(current.wechat)
        await ensure_contact_available(db, target_phone, target_wechat, exclude_lead_id=lead_id)
    assignments = ", ".join(f"{key} = :{key}" for key in values)
    try:
        await db.execute(text(f"UPDATE customer_lead SET {assignments}, updated_at = UTC_TIMESTAMP() WHERE id = :id"), {**values, "id": lead_id})
        await db.execute(text("INSERT INTO business_audit_log (actor_user_id, action, resource_type, resource_id) VALUES (:actor, 'customer_lead.update', 'customer_lead', :id)"), {"actor": account_id, "id": lead_id})
        await db.commit()
    except IntegrityError:
        await db.rollback()
        raise_duplicate_contact()
    return await get_lead(db, lead_id)


async def assign_lead(db: AsyncSession, account_id: int, lead_id: int, request: CustomerLeadAssignment) -> CustomerLead:
    await get_lead(db, lead_id)
    await _validate_owner(db, request)
    await db.execute(text("UPDATE customer_lead SET matchmaker_id = :matchmaker_id, organization_id = :organization_id, updated_at = UTC_TIMESTAMP() WHERE id = :id"), {**request.model_dump(), "id": lead_id})
    await db.execute(text("INSERT INTO business_audit_log (actor_user_id, action, resource_type, resource_id) VALUES (:actor, 'customer_lead.assign', 'customer_lead', :id)"), {"actor": account_id, "id": lead_id})
    await db.commit()
    return await get_lead(db, lead_id)


async def add_follow_up(db: AsyncSession, account_id: int, lead_id: int, request: CustomerLeadFollowUpCreate) -> CustomerLeadFollowUp:
    await get_lead(db, lead_id)
    result = await db.execute(text("""INSERT INTO customer_lead_follow_up
        (lead_id, method, content, intention_level, next_follow_at, created_by)
        VALUES (:lead_id, :method, :content, :intention_level, :next_follow_at, :created_by)"""), {**request.model_dump(), "lead_id": lead_id, "created_by": account_id})
    follow_id = int(result.lastrowid)
    await db.execute(text("UPDATE customer_lead SET status = CASE WHEN status = 'NEW' THEN 'CONTACTED' ELSE status END, next_follow_at = :next_follow_at, updated_at = UTC_TIMESTAMP() WHERE id = :id"), {"next_follow_at": request.next_follow_at, "id": lead_id})
    await db.commit()
    row = (await db.execute(text("SELECT id, lead_id, method, content, intention_level, next_follow_at, created_by, created_at FROM customer_lead_follow_up WHERE id = :id"), {"id": follow_id})).mappings().one()
    return CustomerLeadFollowUp(**dict(row))


async def list_follow_ups(db: AsyncSession, lead_id: int, page: int, page_size: int) -> list[CustomerLeadFollowUp]:
    await get_lead(db, lead_id)
    rows = await db.execute(text("""SELECT id, lead_id, method, content, intention_level, next_follow_at, created_by, created_at
        FROM customer_lead_follow_up WHERE lead_id = :lead_id ORDER BY id DESC LIMIT :limit OFFSET :offset"""), {"lead_id": lead_id, "limit": page_size, "offset": (page - 1) * page_size})
    return [CustomerLeadFollowUp(**dict(row)) for row in rows.mappings().all()]


async def lead_statistics(db: AsyncSession) -> CustomerLeadStatistics:
    row = (await db.execute(text("""SELECT COUNT(*) total, SUM(status = 'NEW') new_count, SUM(status = 'CONTACTED') contacted_count,
        SUM(status = 'INTENDED') intended_count, SUM(status = 'CONVERTED') converted_count, SUM(status = 'LOST') lost_count
        FROM customer_lead"""))).mappings().one()
    return CustomerLeadStatistics(**{key: int(row[key] or 0) for key in ("total", "new_count", "contacted_count", "intended_count", "converted_count", "lost_count")})


async def abandon_lead(db: AsyncSession, account_id: int, lead_id: int, reason: str) -> CustomerLeadAbandonment:
    lead = await get_lead(db, lead_id)
    if lead.status in ("CONVERTED", "CLOSED"):
        raise HTTPException(409, detail="已入库或已关闭的客源不能弃海")
    active = await db.execute(text("SELECT 1 FROM customer_lead_abandonment WHERE lead_id = :id AND restored_at IS NULL"), {"id": lead_id})
    if active.scalar():
        raise HTTPException(409, detail="该客源已在弃海池")
    result = await db.execute(text("""INSERT INTO customer_lead_abandonment (lead_id, reason, abandoned_by)
        VALUES (:lead_id, :reason, :account_id)"""), {"lead_id": lead_id, "reason": reason, "account_id": account_id})
    await db.execute(text("UPDATE customer_lead SET status = 'LOST', matchmaker_id = NULL, next_follow_at = NULL WHERE id = :id"), {"id": lead_id})
    await db.execute(text("INSERT INTO business_audit_log (actor_user_id, action, resource_type, resource_id) VALUES (:actor, 'customer_lead.abandon', 'customer_lead', :id)"), {"actor": account_id, "id": lead_id})
    await db.commit()
    row = (await db.execute(text("SELECT id, lead_id, reason, abandoned_by, abandoned_at, restored_by, restored_at, restore_reason FROM customer_lead_abandonment WHERE id = :id"), {"id": int(result.lastrowid)})).mappings().one()
    return CustomerLeadAbandonment(**dict(row))


async def restore_lead(db: AsyncSession, account_id: int, lead_id: int, reason: str) -> CustomerLead:
    current = await get_lead(db, lead_id)
    active = (await db.execute(text("SELECT id FROM customer_lead_abandonment WHERE lead_id = :id AND restored_at IS NULL ORDER BY id DESC LIMIT 1"), {"id": lead_id})).scalar()
    if not active:
        raise HTTPException(409, detail="该客源不在弃海池")
    # 恢复后线索回到有效状态，先确认联系方式没有被其它有效线索占用。
    await ensure_contact_available(
        db,
        normalize_contact(current.phone),
        normalize_contact(current.wechat),
        exclude_lead_id=lead_id,
        detail="该客源的联系方式已被其他有效线索占用，无法恢复",
    )
    await db.execute(text("UPDATE customer_lead_abandonment SET restored_by = :account_id, restored_at = UTC_TIMESTAMP(), restore_reason = :reason WHERE id = :id"), {"account_id": account_id, "reason": reason, "id": active})
    await db.execute(text("UPDATE customer_lead SET status = 'NEW', updated_at = UTC_TIMESTAMP() WHERE id = :id"), {"id": lead_id})
    await db.execute(text("INSERT INTO business_audit_log (actor_user_id, action, resource_type, resource_id) VALUES (:actor, 'customer_lead.restore', 'customer_lead', :id)"), {"actor": account_id, "id": lead_id})
    await db.commit()
    return await get_lead(db, lead_id)


async def list_abandonments(db: AsyncSession, active_only: bool) -> list[CustomerLeadAbandonment]:
    where = "WHERE restored_at IS NULL" if active_only else ""
    rows = await db.execute(text(f"SELECT id, lead_id, reason, abandoned_by, abandoned_at, restored_by, restored_at, restore_reason FROM customer_lead_abandonment {where} ORDER BY id DESC"))
    return [CustomerLeadAbandonment(**dict(row)) for row in rows.mappings().all()]


async def _insert_leads(
    db: AsyncSession,
    account_id: int,
    rows: list[CustomerLeadImportRow],
    dup_mode: str = "skip",
    *,
    default_source: str = "批量导入",
    default_matchmaker_id: int | None = None,
    default_promoter_id: int | None = None,
    default_audit_status: str = "active",
    default_tags: list[str] | None = None,
    start_index: int = 2,
) -> CustomerLeadImportSummary:
    """批量写入客源线索。留空/无法匹配的字段按 default_* 兜底；dup_mode=skip 时联系方式重复的行跳过。"""
    created = skipped = failed = 0
    errors: list[str] = []
    tag_names = [name.strip() for name in (default_tags or []) if name and name.strip()]
    for index, row in enumerate(rows, start=start_index):  # Excel 数据从第 2 行起
        if not row.phone and not row.wechat:
            failed += 1
            errors.append(f"第 {index} 行：phone 或 wechat 至少提供一个")
            continue
        try:
            if dup_mode == "skip":
                contact_conditions: list[str] = []
                dup_params: dict[str, Any] = {}
                if row.phone:
                    contact_conditions.append("phone = :phone")
                    dup_params["phone"] = row.phone
                if row.wechat:
                    contact_conditions.append("wechat = :wechat")
                    dup_params["wechat"] = row.wechat
                duplicate = await db.execute(
                    text("SELECT id FROM customer_lead WHERE status NOT IN ('LOST', 'CLOSED') AND ("
                         + " OR ".join(contact_conditions) + ") LIMIT 1"),
                    dup_params,
                )
                if duplicate.scalar():
                    skipped += 1
                    continue
            result = await db.execute(text("""INSERT INTO customer_lead
                (name, phone, wechat, source, intention_level, remark,
                 matchmaker_id, promoter_id, audit_status, created_by)
                VALUES (:name, :phone, :wechat, :source, :intention_level, :remark,
                 :matchmaker_id, :promoter_id, :audit_status, :created_by)"""),
                {
                    "name": row.name, "phone": row.phone, "wechat": row.wechat,
                    "source": (row.source or "").strip() or default_source,
                    "intention_level": row.intention_level, "remark": row.remark,
                    "matchmaker_id": default_matchmaker_id, "promoter_id": default_promoter_id,
                    "audit_status": default_audit_status, "created_by": account_id,
                },
            )
            lead_id = int(result.lastrowid)
            await _sync_tags(db, lead_id, tag_names)
            await db.execute(text("""INSERT INTO business_audit_log (actor_user_id, action, resource_type, resource_id)
                VALUES (:actor, 'customer_lead.import', 'customer_lead', :id)"""),
                {"actor": account_id, "id": lead_id},
            )
            created += 1
        except HTTPException:
            raise
        except Exception as exc:  # 单行失败不阻断整批
            failed += 1
            errors.append(f"第 {index} 行：{exc}")
    await db.commit()
    return CustomerLeadImportSummary(created=created, skipped=skipped, failed=failed, total=len(rows), errors=errors[:50])


async def batch_import_leads(
    db: AsyncSession,
    account_id: int,
    rows: list[CustomerLeadImportRow],
    dup_mode: str = "skip",
) -> CustomerLeadBatchImportResult:
    """批量导入客源线索。dup_mode=skip 时联系方式重复的行跳过，append 时仍然导入。"""
    return await _insert_leads(db, account_id, rows, dup_mode)


# ─── Excel 模板下载与文件解析 ────────────────────────────────────────


def build_import_template() -> bytes:
    """生成「客源批量导入」Excel 模板（.xlsx）。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "客源导入"
    sheet.append(_TEMPLATE_HEADERS)
    sheet.append(_TEMPLATE_SAMPLE)
    for index, width in enumerate((18, 18, 18, 14, 12, 36), start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    normalized = str(value).strip()
    return "" if normalized.lower() in {"nan", "none", "null"} else normalized


def _parse_intention(value: Any) -> int:
    raw = _cell_text(value)
    return _INTENTION_ALIASES.get(raw, 1)


def parse_import_file(content: bytes, filename: str) -> list[CustomerLeadImportRow]:
    """解析上传的 .xlsx 客源导入表，返回待写入的行（不做联系方式的即时校验）。"""
    lowered = (filename or "").lower()
    if lowered.endswith(".xls"):
        raise HTTPException(400, detail="暂不支持 .xls 旧格式，请用 Excel 另存为 .xlsx 后重新上传")
    if not lowered.endswith(".xlsx"):
        raise HTTPException(400, detail="仅支持 .xlsx 格式的 Excel 文件")
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # 非 Excel / 文件损坏
        raise HTTPException(400, detail=f"Excel 解析失败：{exc}") from exc
    sheet = workbook.active
    header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
    headers = [_cell_text(cell) for cell in (header_row or ())]

    def find(*keywords: str) -> int | None:
        for index, header in enumerate(headers):
            if any(keyword in header for keyword in keywords):
                return index
        return None

    if find("姓名", "昵称", "name") is None or find("手机", "电话", "phone") is None:
        raise HTTPException(400, detail="模板表头缺少「姓名」或「手机号」列，请使用下载的模板填写")
    name_col, phone_col = find("姓名", "昵称", "name"), find("手机", "电话", "phone")
    wechat_col = find("微信", "wechat")
    source_col = find("来源", "source")
    intention_col = find("意向", "intention")
    remark_col = find("备注", "remark")

    rows: list[CustomerLeadImportRow] = []
    for raw in sheet.iter_rows(min_row=2, values_only=True):
        def pick(column: int | None) -> str:
            if column is None or column >= len(raw):
                return ""
            return _cell_text(raw[column])

        name = pick(name_col)
        phone = pick(phone_col)
        wechat = pick(wechat_col)
        if not name and not phone and not wechat:
            continue  # 整行空白，跳过
        if not name:
            raise HTTPException(400, detail="存在姓名为空的数据行，请补齐后再导入")
        rows.append(
            CustomerLeadImportRow(
                name=name,
                phone=phone or None,
                wechat=wechat or None,
                source=pick(source_col) or "批量导入",
                intention_level=_parse_intention(pick(intention_col)),
                remark=pick(remark_col) or None,
            )
        )
    if not rows:
        raise HTTPException(400, detail="文件中没有可导入的数据行")
    return rows


async def import_leads_file(
    db: AsyncSession,
    account_id: int,
    content: bytes,
    filename: str,
    *,
    audit_status: str = "active",
    default_source: str = "批量导入",
    default_matchmaker_id: int | None = None,
    default_promoter_id: int | None = None,
    default_tags: list[str] | None = None,
    dup_mode: str = "skip",
) -> CustomerLeadImportSummary:
    """解析上传的 Excel 并按导入设置写入客源线索。"""
    rows = parse_import_file(content, filename)
    return await _insert_leads(
        db,
        account_id,
        rows,
        dup_mode,
        default_source=default_source or "批量导入",
        default_matchmaker_id=default_matchmaker_id,
        default_promoter_id=default_promoter_id,
        default_audit_status=audit_status,
        default_tags=default_tags,
    )


async def lead_options(db: AsyncSession) -> CustomerLeadOptions:
    """客源线索页面下拉字典：来源、服务红娘、推广红娘、标签。"""
    sources = [
        CustomerLeadOption(value=row["source"], label=row["source"])
        for row in (await db.execute(text("SELECT DISTINCT source FROM customer_lead WHERE source <> '' ORDER BY source"))).mappings().all()
    ]
    matchmakers = [
        CustomerLeadOption(value=str(row["id"]), label=row["label"])
        for row in (await db.execute(text("""SELECT u.id, COALESCE(NULLIF(u.nickname, ''), u.phone, CONCAT('用户', u.id)) AS label
            FROM user_matchmaker_apply a JOIN users u ON u.id = a.user_id
            JOIN user_role r ON r.user_id = a.user_id AND r.role_code = 'service_matchmaker' AND r.status = 1
            WHERE a.application_type = 'service_matchmaker' AND a.status = 1 ORDER BY u.id"""))).mappings().all()
    ]
    promoters = [
        CustomerLeadOption(value=str(row["id"]), label=row["label"])
        for row in (await db.execute(text("""SELECT u.id, COALESCE(NULLIF(u.nickname, ''), u.phone, CONCAT('用户', u.id)) AS label
            FROM user_matchmaker_apply a JOIN users u ON u.id = a.user_id
            WHERE a.application_type = 'promoter' AND a.status = 1 ORDER BY u.id"""))).mappings().all()
    ]
    tags = [
        CustomerLeadOption(value=row["name"], label=row["name"])
        for row in (await db.execute(text("SELECT name FROM customer_lead_tag WHERE enabled = 1 ORDER BY sort, id"))).mappings().all()
    ]
    return CustomerLeadOptions(sources=sources, matchmakers=matchmakers, promoters=promoters, tags=tags)
