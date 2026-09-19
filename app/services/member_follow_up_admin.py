"""会员 CRM 跟进全览与历史跟进批量导入服务。

对应前端页面：会员CRM → 跟进全览（/love-user-follow-up）与其二级页「导入历史跟进」。
数据源：member_follow_up（会员跟进记录）、users、user_profile、user_auth、matchmaker_admin_account。
"""

from __future__ import annotations

import asyncio
import json
import uuid
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any

import aiofiles
from fastapi import HTTPException, UploadFile
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from PIL import Image, ImageOps, UnidentifiedImageError
from PIL.Image import DecompressionBombError
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import settings
from app.schemas.member_follow_up_admin import (
    MemberFollowUp,
    MemberFollowUpImportResult,
    MemberFollowUpListPage,
    MemberFollowUpPage,
    MemberFollowUpRow,
    MemberFollowUpSummary,
)

# 跟进方式库内枚举 -> 中文
_METHOD_LABELS: dict[str, str] = {
    "PHONE": "电话",
    "WECHAT": "微信",
    "VISIT": "面谈",
    "OTHER": "其他",
}

# 导入模板表头（顺序即模板列顺序）
_TEMPLATE_HEADERS: list[str] = ["会员手机号", "跟进内容", "跟进人", "跟进时间", "跟进方式"]

# 模板示例行
_TEMPLATE_SAMPLE: list[str] = [
    "13800000000",
    "示例：电话沟通，客户反馈本周有空到店，已约定周六下午",
    "admin",
    "2026-01-01 10:00:00",
    "电话",
]

# 跟进方式中文 -> 库内枚举
_METHOD_ALIASES: dict[str, str] = {
    "电话": "PHONE",
    "phone": "PHONE",
    "微信": "WECHAT",
    "wechat": "WECHAT",
    "vx": "WECHAT",
    "面谈": "VISIT",
    "拜访": "VISIT",
    "到店": "VISIT",
    "visit": "VISIT",
    "其他": "OTHER",
    "other": "OTHER",
}

# 跟进时间可接受的文本格式
_TIME_FORMATS: tuple[str, ...] = (
    "%Y-%m-%d %H:%M:%S",
    "%Y-%m-%d %H:%M",
    "%Y-%m-%d",
    "%Y/%m/%d %H:%M:%S",
    "%Y/%m/%d %H:%M",
    "%Y/%m/%d",
    "%Y.%m.%d %H:%M:%S",
    "%Y.%m.%d",
)

# 单次导入行数上限，避免超大文件打爆请求
_MAX_IMPORT_ROWS = 2000

_LIST_BODY = """
SELECT f.id, f.user_id, f.method, f.content, f.created_at,
       u.nickname, u.avatar,
       CONCAT('G', LPAD(u.id, 6, '0')) AS member_code,
       p.intention_level,
       COALESCE(acc.display_name, mk.nickname, 'admin') AS matchmaker_name
FROM member_follow_up f
JOIN users u ON u.id = f.user_id
LEFT JOIN user_profile p ON p.user_id = u.id
LEFT JOIN user_auth au ON au.user_id = u.id
LEFT JOIN matchmaker_admin_account acc ON acc.id = f.created_by
LEFT JOIN users mk ON mk.id = f.created_by
"""

# 计数用（不需要展示列，避免多算 join）
_COUNT_BODY = """
FROM member_follow_up f
JOIN users u ON u.id = f.user_id
LEFT JOIN user_profile p ON p.user_id = u.id
LEFT JOIN user_auth au ON au.user_id = u.id
"""

# 8 个时间维度计数：与前端「跟进全览」统计卡顺序一致
# 本周/上周按自然周（周一为起点）计算，避免跨年时 YEARWEEK 边界问题
_SUMMARY_SQL = """
SELECT
    COUNT(*) AS all_count,
    COALESCE(SUM(DATE(created_at) = CURDATE()), 0) AS today_count,
    COALESCE(SUM(DATE(created_at) = DATE_SUB(CURDATE(), INTERVAL 1 DAY)), 0) AS yesterday_count,
    COALESCE(SUM(DATE(created_at) >= DATE_SUB(CURDATE(), INTERVAL 2 DAY)), 0) AS three_days_count,
    COALESCE(SUM(DATE(created_at) >= DATE_SUB(CURDATE(), INTERVAL WEEKDAY(CURDATE()) DAY)), 0)
        AS this_week_count,
    COALESCE(SUM(DATE(created_at) >= DATE_SUB(CURDATE(), INTERVAL WEEKDAY(CURDATE()) + 7 DAY)
        AND DATE(created_at) < DATE_SUB(CURDATE(), INTERVAL WEEKDAY(CURDATE()) DAY)), 0)
        AS last_week_count,
    COALESCE(SUM(DATE_FORMAT(created_at, '%Y-%m') = DATE_FORMAT(CURDATE(), '%Y-%m')), 0)
        AS this_month_count,
    COALESCE(SUM(DATE_FORMAT(created_at, '%Y-%m')
        = DATE_FORMAT(DATE_SUB(CURDATE(), INTERVAL 1 MONTH), '%Y-%m')), 0)
        AS last_month_count
FROM member_follow_up
"""


def _to_int(value: Any) -> int:
    return int(value or 0)


def _method_label(method: str | None) -> str:
    return _METHOD_LABELS.get(str(method or "").upper(), "其他")


def _row(payload: dict[str, Any]) -> MemberFollowUpRow:
    """数据库行 -> 列表行。"""
    method = str(payload.get("method") or "OTHER")
    intention = payload.get("intention_level")
    return MemberFollowUpRow(
        id=int(payload["id"]),
        user_id=int(payload["user_id"]),
        member_code=payload.get("member_code") or f"G{int(payload['user_id']):06d}",
        nickname=payload.get("nickname"),
        avatar=payload.get("avatar"),
        matchmaker_name=payload.get("matchmaker_name") or "admin",
        method=method,
        method_label=_method_label(method),
        content=payload.get("content") or "",
        # member_follow_up 无「是否系统自动生成」标记列，数据源不支持，恒为 null
        note=None,
        intention_level=int(intention) if intention is not None else None,
        created_at=payload.get("created_at"),
    )


async def list_follow_ups(
    db: AsyncSession,
    page: int,
    page_size: int,
    keyword: str | None = None,
    intention_level: int | None = None,
    start_date: str | None = None,
    end_date: str | None = None,
) -> MemberFollowUpListPage:
    """跟进全览分页列表。keyword 匹配昵称/手机号/实名/会员编号。"""
    conditions = ["1=1"]
    params: dict[str, Any] = {"limit": page_size, "offset": (page - 1) * page_size}
    if keyword:
        conditions.append(
            "(u.nickname LIKE CONCAT('%', :keyword, '%')"
            " OR u.phone LIKE CONCAT('%', :keyword, '%')"
            " OR au.real_name LIKE CONCAT('%', :keyword, '%')"
            " OR CONCAT('G', LPAD(u.id, 6, '0')) LIKE CONCAT('%', :keyword, '%'))"
        )
        params["keyword"] = keyword
    if intention_level is not None:
        conditions.append("p.intention_level = :intention_level")
        params["intention_level"] = intention_level
    if start_date:
        conditions.append("f.created_at >= :start_date")
        params["start_date"] = start_date
    if end_date:
        conditions.append("f.created_at < DATE_ADD(:end_date, INTERVAL 1 DAY)")
        params["end_date"] = end_date
    where = " AND ".join(conditions)
    rows = await db.execute(
        text(
            f"{_LIST_BODY} WHERE {where}"
            " ORDER BY f.created_at DESC, f.id DESC LIMIT :limit OFFSET :offset"
        ),
        params,
    )
    count_params = {k: v for k, v in params.items() if k not in ("limit", "offset")}
    total = _to_int(
        (await db.scalar(text(f"SELECT COUNT(*) {_COUNT_BODY} WHERE {where}"), count_params))
    )
    return MemberFollowUpListPage(
        items=[_row(dict(row)) for row in rows.mappings().all()],
        page=page,
        page_size=page_size,
        total=total,
        has_more=page * page_size < total,
    )


async def follow_up_summary(db: AsyncSession) -> MemberFollowUpSummary:
    """跟进全览 8 个时间统计卡计数。"""
    row = (await db.execute(text(_SUMMARY_SQL))).mappings().one()
    return MemberFollowUpSummary(
        all=_to_int(row["all_count"]),
        today=_to_int(row["today_count"]),
        yesterday=_to_int(row["yesterday_count"]),
        three_days=_to_int(row["three_days_count"]),
        this_week=_to_int(row["this_week_count"]),
        last_week=_to_int(row["last_week_count"]),
        this_month=_to_int(row["this_month_count"]),
        last_month=_to_int(row["last_month_count"]),
    )


def build_import_template() -> bytes:
    """生成「历史跟进」导入 Excel 模板（.xlsx）。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "历史跟进"
    sheet.append(_TEMPLATE_HEADERS)
    sheet.append(_TEMPLATE_SAMPLE)
    for index, width in enumerate((18, 46, 16, 22, 12), start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    normalized = str(value).strip()
    return "" if normalized.lower() in {"nan", "none", "null"} else normalized


def _parse_datetime(value: Any) -> datetime | None:
    """解析跟进时间；Excel 原生日期时间单元格直接取用，文本按多种格式尝试。"""
    if isinstance(value, datetime):
        return value
    raw = _cell_text(value)
    if not raw:
        return None
    for fmt in _TIME_FORMATS:
        try:
            return datetime.strptime(raw, fmt)
        except ValueError:
            continue
    return None


def _parse_method(value: Any) -> str:
    """解析跟进方式，无法识别统一按 OTHER。"""
    raw = _cell_text(value)
    if not raw:
        return "OTHER"
    return _METHOD_ALIASES.get(raw, _METHOD_ALIASES.get(raw.lower(), "OTHER"))


def _find_column(headers: list[str], *keywords: str) -> int | None:
    """按关键字定位模板列下标（表头可能被用户微调，故用包含匹配）。"""
    for index, header in enumerate(headers):
        for keyword in keywords:
            if keyword in header:
                return index
    return None


def _read_sheet(data: bytes, filename: str) -> tuple[list[str], list[tuple[int, list[Any]]]]:
    """解析上传文件，返回 (表头列表, [(Excel 行号, 行值列表)])。仅支持 .xlsx。"""
    lowered = (filename or "").lower()
    if lowered.endswith(".xls"):
        raise HTTPException(400, detail="暂不支持 .xls 旧格式，请用 Excel 另存为 .xlsx 后重新上传")
    if not lowered.endswith(".xlsx"):
        raise HTTPException(400, detail="仅支持 .xlsx 格式的 Excel 文件")
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # 非 Excel / 文件损坏
        raise HTTPException(400, detail=f"Excel 解析失败：{exc}") from exc
    try:
        sheet = workbook.active
        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
        headers = [_cell_text(cell) for cell in (header_row or ())]
        if _find_column(headers, "手机", "电话", "phone") is None:
            raise HTTPException(400, detail="模板表头缺少「会员手机号」列，请使用下载的模板填写")
        if _find_column(headers, "跟进内容", "内容", "content") is None:
            raise HTTPException(400, detail="模板表头缺少「跟进内容」列，请使用下载的模板填写")
        rows: list[tuple[int, list[Any]]] = []
        for index, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if len(rows) >= _MAX_IMPORT_ROWS:
                raise HTTPException(400, detail=f"单次最多导入 {_MAX_IMPORT_ROWS} 行，请拆分文件后重试")
            rows.append((index, list(values or ())))
        return headers, rows
    finally:
        workbook.close()


async def _default_operator_id(db: AsyncSession, account_id: int) -> int:
    """导入时「跟进人」取不到时的默认账号：优先平台 admin 账号，回退当前操作人。"""
    admin_id = await db.scalar(
        text("SELECT id FROM matchmaker_admin_account WHERE username = 'admin' LIMIT 1")
    )
    return int(admin_id or account_id)


async def _resolve_operator_id(
    db: AsyncSession, name: str, cache: dict[str, int], fallback: int
) -> int:
    """把模板中的「跟进人」称呼映射到后台账号 id；匹配不到（或为空）回退默认账号。"""
    key = name.strip()
    if not key:
        return fallback
    if key in cache:
        return cache[key]
    operator_id = await db.scalar(
        text(
            "SELECT id FROM matchmaker_admin_account"
            " WHERE display_name = :name AND status = 1 LIMIT 1"
        ),
        {"name": key},
    )
    if not operator_id:
        # 服务红娘可能只有 users.nickname，通过账号绑定的 matchmaker_user_id 反查
        operator_id = await db.scalar(
            text(
                "SELECT acc.id FROM users u"
                " JOIN matchmaker_admin_account acc ON acc.matchmaker_user_id = u.id"
                " WHERE u.nickname = :name AND acc.status = 1 LIMIT 1"
            ),
            {"name": key},
        )
    resolved = int(operator_id or fallback)
    cache[key] = resolved
    return resolved


async def import_follow_ups(
    db: AsyncSession,
    account_id: int,
    data: bytes,
    filename: str,
) -> MemberFollowUpImportResult:
    """解析上传的 Excel，按手机号匹配会员后批量写入历史跟进记录。

    规则（与页面「须知」一致）：
    - 按「会员手机号」识别会员，会员CRM 中不存在该手机号的行无法导入；
    - 「跟进人」与后台账号称呼无法对应或为空时默认 admin；
    - 「跟进时间」缺失或格式无法匹配时使用导入时间；
    - 「跟进方式」无法识别时按 OTHER。
    """
    if not data:
        raise HTTPException(400, detail="上传文件为空")

    headers, rows_only = _read_sheet(data, filename)
    if not rows_only:
        # 表头存在但无数据行，视为空文件
        return MemberFollowUpImportResult(created=0, skipped=0, failed=0, errors=[])

    phone_col = _find_column(headers, "手机", "电话", "phone")
    content_col = _find_column(headers, "跟进内容", "内容", "content")
    matchmaker_col = _find_column(headers, "跟进人", "红娘", "matchmaker")
    time_col = _find_column(headers, "跟进时间", "时间", "time")
    method_col = _find_column(headers, "跟进方式", "方式", "method")

    fallback_operator = await _default_operator_id(db, account_id)
    operator_cache: dict[str, int] = {}

    created = skipped = failed = 0
    errors: list[str] = []
    for excel_row, values in rows_only:

        def cell(column: int | None) -> Any:
            return values[column] if column is not None and column < len(values) else None

        phone = _cell_text(cell(phone_col))
        content = _cell_text(cell(content_col))
        if not phone and not content:
            skipped += 1
            continue
        if not phone:
            failed += 1
            errors.append(f"第 {excel_row} 行：缺少会员手机号")
            continue
        if not content:
            failed += 1
            errors.append(f"第 {excel_row} 行：缺少跟进内容")
            continue

        user_id = await db.scalar(
            text("SELECT id FROM users WHERE phone = :phone AND status = 1 LIMIT 1"),
            {"phone": phone},
        )
        if not user_id:
            failed += 1
            errors.append(f"第 {excel_row} 行：手机号 {phone} 在会员CRM中不存在，无法导入")
            continue

        method = _parse_method(cell(method_col))
        followed_at = _parse_datetime(cell(time_col))
        operator_id = await _resolve_operator_id(
            db, _cell_text(cell(matchmaker_col)), operator_cache, fallback_operator
        )
        try:
            if followed_at is None:
                result = await db.execute(
                    text(
                        "INSERT INTO member_follow_up (user_id, method, content, created_by)"
                        " VALUES (:user_id, :method, :content, :created_by)"
                    ),
                    {
                        "user_id": int(user_id),
                        "method": method,
                        "content": content,
                        "created_by": operator_id,
                    },
                )
            else:
                result = await db.execute(
                    text(
                        "INSERT INTO member_follow_up (user_id, method, content, created_by, created_at)"
                        " VALUES (:user_id, :method, :content, :created_by, :created_at)"
                    ),
                    {
                        "user_id": int(user_id),
                        "method": method,
                        "content": content,
                        "created_by": operator_id,
                        "created_at": followed_at,
                    },
                )
            follow_id = int(result.lastrowid)
            await db.execute(
                text(
                    "INSERT INTO business_audit_log (actor_user_id, action, resource_type, resource_id)"
                    " VALUES (:actor, 'member.follow_up.import', 'member_follow_up', :id)"
                ),
                {"actor": account_id, "id": follow_id},
            )
            created += 1
        except HTTPException:
            raise
        except Exception as exc:  # 单行失败不阻断整批
            failed += 1
            errors.append(f"第 {excel_row} 行：{exc}")
    await db.commit()
    return MemberFollowUpImportResult(
        created=created, skipped=skipped, failed=failed, errors=errors[:50]
    )


# ─── 新增跟进（文字 + 图片 + 录音） ─────────────────────────────

FOLLOW_UP_MAX_IMAGES = 9
FOLLOW_UP_IMAGE_MAX_BYTES = 5 * 1024 * 1024
FOLLOW_UP_VOICE_MAX_BYTES = 20 * 1024 * 1024
_FOLLOW_UP_IMAGE_MAX_PIXELS = 25_000_000
_FOLLOW_UP_VOICE_EXT: dict[str, str] = {
    "audio/mpeg": "mp3",
    "audio/mp3": "mp3",
    "audio/wav": "wav",
    "audio/x-wav": "wav",
    "audio/wave": "wav",
    "audio/mp4": "m4a",
    "audio/x-m4a": "m4a",
    "audio/aac": "aac",
    "audio/ogg": "ogg",
    "audio/amr": "amr",
    "audio/webm": "webm",
}


async def _fu_read_limited(file: UploadFile, limit: int) -> bytes:
    chunks: list[bytes] = []
    total = 0
    while True:
        chunk = await file.read(1024 * 1024)
        if not chunk:
            break
        total += len(chunk)
        if total > limit:
            raise HTTPException(413, detail=f"文件大小不能超过{limit // 1024 // 1024}MB")
        chunks.append(chunk)
    return b"".join(chunks)


async def _fu_write_bytes(path: Path, data: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    async with aiofiles.open(path, "wb") as output:
        await output.write(data)


def _fu_image_outputs(data: bytes) -> tuple[bytes, bytes]:
    """图片 -> (webp 原图, webp 缩略图)，口径与会员资料上传一致。"""
    try:
        with Image.open(BytesIO(data)) as source:
            if source.format not in {"JPEG", "PNG"}:
                raise HTTPException(415, detail="仅支持JPG、JPEG或PNG图片")
            if source.width * source.height > _FOLLOW_UP_IMAGE_MAX_PIXELS:
                raise HTTPException(413, detail="图片像素不能超过2500万")
            source.verify()
        with Image.open(BytesIO(data)) as source:
            image = ImageOps.exif_transpose(source)
            if image.mode not in {"RGB", "RGBA"}:
                image = image.convert("RGBA" if "A" in image.getbands() else "RGB")
            output = BytesIO()
            image.save(output, format="WEBP", quality=85, method=4)
            thumbnail = image.copy()
            thumbnail.thumbnail((480, 480), Image.Resampling.LANCZOS)
            thumb_output = BytesIO()
            thumbnail.save(thumb_output, format="WEBP", quality=80, method=4)
            return output.getvalue(), thumb_output.getvalue()
    except DecompressionBombError as exc:
        raise HTTPException(413, detail="图片像素过大") from exc
    except (UnidentifiedImageError, OSError) as exc:
        raise HTTPException(415, detail="图片内容无法识别") from exc
    except HTTPException:
        raise


def _voice_suffix(voice: UploadFile) -> str:
    """按 Content-Type / 文件名后缀推断录音扩展名，白名单校验。"""
    mime = (voice.content_type or "").split(";")[0].strip().lower()
    if mime in _FOLLOW_UP_VOICE_EXT:
        return _FOLLOW_UP_VOICE_EXT[mime]
    name_ext = (voice.filename or "").rsplit(".", 1)[-1].strip().lower()
    if name_ext in set(_FOLLOW_UP_VOICE_EXT.values()):
        return name_ext
    raise HTTPException(415, detail="仅支持 mp3/wav/m4a/aac/ogg/amr/webm 录音")


def _parse_follow_up_images(value: Any) -> list[str]:
    if value is None:
        return []
    if isinstance(value, list):
        return [str(item) for item in value]
    try:
        parsed = json.loads(value)
    except (TypeError, ValueError):
        return []
    return [str(item) for item in parsed] if isinstance(parsed, list) else []


def _build_follow_up_row(row: Any) -> MemberFollowUp:
    return MemberFollowUp(
        id=int(row["id"]),
        user_id=int(row["user_id"]),
        method=str(row["method"]),
        content=str(row["content"] or ""),
        next_follow_at=row["next_follow_at"],
        created_by=int(row["created_by"]),
        created_at=row["created_at"],
        images=_parse_follow_up_images(row["images"]),
        voice_url=row["voice_url"],
        voice_duration_sec=int(row["voice_duration_sec"]) if row["voice_duration_sec"] is not None else None,
        matchmaker_name=row["matchmaker_name"],
    )


_FOLLOW_UP_SELECT = """
SELECT f.id, f.user_id, f.method, f.content, f.next_follow_at, f.created_by, f.created_at,
       f.images, f.voice_url, f.voice_duration_sec,
       COALESCE(acc.display_name, mk.nickname) AS matchmaker_name
FROM member_follow_up f
LEFT JOIN matchmaker_admin_account acc ON acc.id = f.created_by
LEFT JOIN users mk ON mk.id = f.created_by
"""


async def get_member_follow_ups(
    db: AsyncSession, member_id: int, page: int, page_size: int
) -> MemberFollowUpPage:
    """查询单个会员的跟进记录（含图片/录音与跟进红娘称呼），倒序分页。"""
    if not await db.scalar(text("SELECT 1 FROM users WHERE id = :id"), {"id": member_id}):
        raise HTTPException(404, detail="会员不存在")
    params: dict[str, Any] = {"uid": member_id, "limit": page_size, "offset": (page - 1) * page_size}
    rows = await db.execute(
        text(f"{_FOLLOW_UP_SELECT} WHERE f.user_id = :uid ORDER BY f.id DESC LIMIT :limit OFFSET :offset"),
        params,
    )
    total = int(
        (await db.scalar(text("SELECT COUNT(*) FROM member_follow_up WHERE user_id = :uid"), {"uid": member_id})) or 0
    )
    items = [_build_follow_up_row(row) for row in rows.mappings().all()]
    return MemberFollowUpPage(items=items, page=page, page_size=page_size, total=total, has_more=page * page_size < total)


async def create_follow_up_with_media(
    db: AsyncSession,
    member_id: int,
    method: str,
    content: str,
    next_follow_at: datetime | None,
    matchmaker_id: int | None,
    images: list[UploadFile] | None,
    voice: UploadFile | None,
    voice_duration_sec: int | None,
    actor_id: int,
) -> MemberFollowUp:
    """新增跟进记录，支持文字 + 图片（转 webp，最多 9 张）+ 录音。

    ``matchmaker_id`` 为空时记录当前操作账号；提供时必须是启用中的后台账号。
    图片存 ``{upload_dir}/{member_id}/follow-up-img-{uuid}.webp``，录音存
    ``follow-up-voice-{uuid}.{ext}``；URL 形如 ``/storage/uploads/{member_id}/...``。
    """
    if not await db.scalar(text("SELECT 1 FROM users WHERE id = :id"), {"id": member_id}):
        raise HTTPException(404, detail="会员不存在")

    normalized_content = (content or "").strip()
    image_files = [item for item in (images or []) if item is not None and (item.filename or "").strip()]
    has_voice = voice is not None and bool((voice.filename or "").strip()) and voice.size != 0
    if not normalized_content and not image_files and not has_voice:
        raise HTTPException(422, detail="跟进文字与图片/录音至少提供一项")

    if matchmaker_id is None:
        operator_id = actor_id
    else:
        operator_id = await db.scalar(
            text("SELECT id FROM matchmaker_admin_account WHERE id = :id AND status = 1"),
            {"id": matchmaker_id},
        )
        if not operator_id:
            raise HTTPException(422, detail="跟进红娘不存在或已停用")
        operator_id = int(operator_id)

    if len(image_files) > FOLLOW_UP_MAX_IMAGES:
        raise HTTPException(422, detail=f"图片最多{FOLLOW_UP_MAX_IMAGES}张")

    directory = Path(settings.upload_dir) / str(member_id)
    image_urls: list[str] = []
    for image in image_files:
        data = await _fu_read_limited(image, FOLLOW_UP_IMAGE_MAX_BYTES)
        webp_data, _thumb = await asyncio.to_thread(_fu_image_outputs, data)
        name = uuid.uuid4().hex
        target = directory / f"follow-up-img-{name}.webp"
        await _fu_write_bytes(target, webp_data)
        image_urls.append(f"/storage/uploads/{member_id}/{target.name}")

    voice_url: str | None = None
    if has_voice:
        voice_data = await _fu_read_limited(voice, FOLLOW_UP_VOICE_MAX_BYTES)  # type: ignore[union-attr]
        suffix = _voice_suffix(voice)  # type: ignore[arg-type]
        name = uuid.uuid4().hex
        target = directory / f"follow-up-voice-{name}.{suffix}"
        await _fu_write_bytes(target, voice_data)
        voice_url = f"/storage/uploads/{member_id}/{target.name}"

    result = await db.execute(
        text(
            "INSERT INTO member_follow_up (user_id, method, content, next_follow_at, created_by, images, voice_url, voice_duration_sec)"
            " VALUES (:uid, :method, :content, :next_follow_at, :created_by, :images, :voice_url, :voice_duration_sec)"
        ),
        {
            "uid": member_id,
            "method": method,
            "content": normalized_content,
            "next_follow_at": next_follow_at,
            "created_by": operator_id,
            "images": json.dumps(image_urls, ensure_ascii=False) if image_urls else None,
            "voice_url": voice_url,
            "voice_duration_sec": voice_duration_sec if (voice_url and voice_duration_sec) else None,
        },
    )
    follow_id = int(result.lastrowid)
    await db.execute(
        text(
            "INSERT INTO business_audit_log (actor_user_id, action, resource_type, resource_id, after_json)"
            " VALUES (:actor, 'member.follow_up.create', 'member_follow_up', :rid, :after)"
        ),
        {
            "actor": actor_id,
            "rid": follow_id,
            "after": json.dumps(
                {"images": image_urls, "voice_url": voice_url, "created_by": operator_id},
                ensure_ascii=False,
            ),
        },
    )
    await db.commit()
    row = (
        await db.execute(text(f"{_FOLLOW_UP_SELECT} WHERE f.id = :id"), {"id": follow_id})
    ).mappings().one()
    return _build_follow_up_row(row)
